import os
from pathlib import Path
from groq import Groq
from dotenv import load_dotenv
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document

load_dotenv()
myapikey=os.getenv("GROQ_API_KEY")
if not myapikey:
    raise ValueError("No API Key")
client=Groq(api_key=myapikey)
model="llama-3.3-70b-versatile"
job_description="""
Description
Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.
Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.
Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older
Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills
"""

class JobDes(BaseModel):
    role:str
    experience:float|None
    required_skills: list[str]
    preffered_skills:list[str]
    educational_requirements:list[str]
    responsibilities:list[str]
jobdes_schema=JobDes.model_json_schema()

prompt_system=f""" You are an expert HR assistant. Your job is to read the job description carefully
                and extract all useful information.
                Follow this scema {jobdes_schema}. Do not return the schema itself. 
                Do not invent new information. 
                Return the answer in JSON."""
prompt_user=f"""Read the job description  {job_description} and analyse the information."""
message_system={"role":"system", "content":prompt_system}
message={"role":"user","content":prompt_user}
response_format={"type":"json_object"}
messages=[message_system,message]
response=client.chat.completions.create(model=model,messages=messages,temperature=0,response_format=response_format)
answer= response.choices[0].message.content
#print(answer)

import json
jobdata=json.loads(answer)
job=JobDes(**jobdata)
print(job.experience)
print(job.educational_requirements)

class Experience(BaseModel):
    company_name:str |None=None
    role:str|None=None
    duration:str|None=None
    description:str|None=None
class Resume(BaseModel):
    name:str |None=None
    phone:str |None=None
    mail:str |None=None
    total_experience_years:float |None=None
    skills:list[str] =[]
    experience:list[Experience]=[]
    education:list[str]=[]
    projects:list[str]=[]
class MatchResult(BaseModel):
    score:float
    details:dict


resume_schema=Resume.model_json_schema()

def final_score(job,resume):
    match_schema=MatchResult.model_json_schema()
    prompt=f"""You are an HR recruiter
    Compare the candidat's resume with the job description
    Job description: {job.model_dump_json(indent=2)}
    Candidate's Resume: {resume.model_dump_json(indent=2)}

    Return the result matching this schema
    {match_schema}
     Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the result concise and easy to read and resturn result in JSON
    """
    message={"role":"user","content":prompt}
    messages=[message]
    response_format={"type":"json_object"}
    response=client.chat.completions.create(model=model,messages=messages,temperature=0,response_format=response_format)
    data=json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
     Experience
     Professional Experience
     Work History
     Employment
     Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt=f"""
    Analyze the following resume {resume_text}"""
    
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    rjson=response.choices[0].message.content
    data = json.loads(rjson)
    resume = Resume(**data)
    return resume


def read_pdf(file_path):
    reader=PdfReader(file_path)
    text=""
    for page in reader.pages:
        p=page.extract_text()
        if p:
            text+=p + "\n"
    return text
def read_document(file_path):
    text=""
    document=Document(file_path)
    for i in document.paragraphs:
        if i.text.strip():
            text+=i.text +"\n"
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text
def read_resume(file_path):
    if file_path.suffix.lower()==".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower()==".docx":
        return read_document(file_path)
    else:
        return None
resume_folder=Path(r"C:\Users\Sumit\course\week1\miniproj\Resume_folder")
all_result=[]
import time
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf",".docx"]:
        continue
    print("\nProcessing:  ", file_path.name)
    resume_text=read_resume(file_path)
    parsed_resume=parse_resume(resume_text)
    time.sleep(5)
    result=final_score(job,parsed_resume)
    time.sleep(5)
    print(f"Score: {result.score}")
    all_result.append({"Name":parsed_resume.name,"Score":result.score,"Details":result.details})
    all_result.sort(key=lambda candidate: candidate["Score"],reverse=True)
    top_2=all_result[:2]
    bot_2=all_result[-2:]
    print("Top 2")
    for i in top_2:
        print(i["Name"],"-", i["Score"],"%")
        print(i["Details"])
    print("Bottom 2")
    for i in bot_2:
        print(i["Name"],"-", i["Score"],"%")
        print(i["Details"])